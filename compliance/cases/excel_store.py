"""Excel-backed case register.

The workbook (`data/cases/cases.xlsx`) is meant to be opened and edited by
compliance staff directly, so it stays deliberately simple:

  Disclaimer information-only notice (first sheet, EN/AR)
  Cases      one row per open-ended compliance question / AI use case
  Documents  one row per required document of a case
  Log        append-only audit trail of changes

Writes go through a process-wide lock and an atomic file replace. Uploaded
files are stored under `data/cases/<case_id>/`.
"""
from __future__ import annotations

import json
import os
import re
import secrets
import threading
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from compliance import disclaimer
from compliance.config import get_settings

CASE_COLUMNS = [
    "case_id", "created_at", "updated_at", "owner", "jurisdictions", "sector", "ai_use_case",
    "question", "status", "answer_status", "related_provisions", "reviewer", "review_note", "notes",
]
DOC_COLUMNS = [
    "case_id", "doc_id", "name", "kind", "mandatory", "basis", "status", "due_date", "file", "updated_at", "note",
]
LOG_COLUMNS = ["timestamp", "case_id", "actor", "action", "detail"]
DISCLAIMER_SHEET = "Disclaimer"

CASE_STATUSES = ("open", "in_review", "approved", "closed")
DOC_STATUSES = ("missing", "received", "waived", "not_applicable")

_CASE_ID = re.compile(r"^CASE-[0-9A-F]{8}$")
_SAFE_NAME = re.compile(r"[^A-Za-z0-9._-]+")
_lock = threading.RLock()


class CaseError(ValueError):
    pass


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _check_case_id(case_id: str) -> str:
    if not _CASE_ID.match(case_id or ""):
        raise CaseError(f"Invalid case id: {case_id!r}")
    return case_id


class CaseStore:
    def __init__(self, path: Path | None = None):
        self.path = path or get_settings().cases_workbook
        self.files_dir = self.path.parent

    # ---------------- workbook helpers ----------------
    def _open(self):
        if not self.path.exists():
            self.path.parent.mkdir(parents=True, exist_ok=True)
            wb = Workbook()
            wb.active.title = "Cases"
            for title, cols in (("Cases", CASE_COLUMNS), ("Documents", DOC_COLUMNS), ("Log", LOG_COLUMNS)):
                ws = wb[title] if title in wb.sheetnames else wb.create_sheet(title)
                ws.append(cols)
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                ws.freeze_panes = "A2"
            self._add_disclaimer_sheet(wb)
            self._save(wb)
        wb = load_workbook(self.path)
        if DISCLAIMER_SHEET not in wb.sheetnames:  # workbooks created by earlier versions
            self._add_disclaimer_sheet(wb)
            self._save(wb)
        return wb

    @staticmethod
    def _add_disclaimer_sheet(wb) -> None:
        """The workbook may be shared on its own, so it carries the disclaimer itself."""
        ws = wb.create_sheet(DISCLAIMER_SHEET, 0)
        ws.column_dimensions["A"].width = 120
        ws.append(["For information only - not legal advice. All legal responsibility rests with the user."])
        ws["A1"].font = Font(bold=True, size=13)
        for lang in ("en", "ar"):
            ws.append([""])
            ws.append([disclaimer.full(lang)])
            ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        wb.active = 0

    def _save(self, wb) -> None:
        tmp = self.path.with_suffix(".tmp.xlsx")
        wb.save(tmp)
        os.replace(tmp, self.path)

    @staticmethod
    def _rows(ws: Worksheet) -> list[dict]:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [str(h) for h in rows[0]]
        return [
            {k: ("" if v is None else v) for k, v in zip(header, r)}
            for r in rows[1:]
            if any(v not in (None, "") for v in r)
        ]

    @staticmethod
    def _find_row(ws: Worksheet, matches) -> int | None:
        header = [c.value for c in ws[1]]
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if matches(dict(zip(header, row))):
                return idx
        return None

    @staticmethod
    def _set(ws: Worksheet, row: int, column: str, value) -> None:
        header = [c.value for c in ws[1]]
        ws.cell(row=row, column=header.index(column) + 1, value=value)

    def _log(self, wb, case_id: str, actor: str, action: str, detail: str = "") -> None:
        wb["Log"].append([_now(), case_id, actor or "system", action, detail])

    # ---------------- cases ----------------
    def create_case(
        self,
        owner: str,
        question: str,
        jurisdictions: list[str],
        sector: str = "",
        ai_use_case: str = "",
        notes: str = "",
        answer: dict | None = None,
    ) -> str:
        if not question.strip():
            raise CaseError("A case needs a question")
        with _lock:
            wb = self._open()
            case_id = "CASE-" + secrets.token_hex(4).upper()
            provisions = []
            answer_status = ""
            if answer:
                answer_status = answer.get("status", "")
                for c in answer.get("findings", []):
                    provisions.append(f"{c.get('source_id')} art. {c.get('article')}")
            row = {
                "case_id": case_id,
                "created_at": _now(),
                "updated_at": _now(),
                "owner": owner,
                "jurisdictions": ", ".join(jurisdictions),
                "sector": sector,
                "ai_use_case": ai_use_case,
                "question": question,
                "status": "open",
                "answer_status": answer_status,
                "related_provisions": "; ".join(dict.fromkeys(provisions)),
                "reviewer": "",
                "review_note": "",
                "notes": notes,
            }
            wb["Cases"].append([row[c] for c in CASE_COLUMNS])
            for doc in (answer or {}).get("documents", []):
                self._append_document(wb, case_id, doc)
            self._log(wb, case_id, owner, "created", question[:200])
            self._save(wb)
        (self.files_dir / case_id).mkdir(parents=True, exist_ok=True)
        return case_id

    def list_cases(self, status: str | None = None) -> list[dict]:
        with _lock:
            wb = self._open()
            cases = self._rows(wb["Cases"])
            docs = self._rows(wb["Documents"])
        for case in cases:
            case["missing_mandatory"] = sum(
                1 for d in docs if d["case_id"] == case["case_id"] and _truthy(d["mandatory"]) and d["status"] == "missing"
            )
        if status:
            cases = [c for c in cases if c["status"] == status]
        return cases

    def get_case(self, case_id: str) -> dict:
        _check_case_id(case_id)
        with _lock:
            wb = self._open()
            case = next((c for c in self._rows(wb["Cases"]) if c["case_id"] == case_id), None)
            if case is None:
                raise CaseError(f"Case {case_id} not found")
            docs = [d for d in self._rows(wb["Documents"]) if d["case_id"] == case_id]
            log = [entry for entry in self._rows(wb["Log"]) if entry["case_id"] == case_id]
        for d in docs:
            d["mandatory"] = _truthy(d["mandatory"])
            d["basis"] = _json_or_empty(d["basis"])
        case["documents"] = docs
        case["alerts"] = _alerts(docs)
        case["log"] = log
        return case

    def update_case(self, case_id: str, actor: str, **fields) -> dict:
        _check_case_id(case_id)
        allowed = {"status", "reviewer", "review_note", "notes", "owner", "sector"}
        unknown = set(fields) - allowed
        if unknown:
            raise CaseError(f"Cannot update fields: {sorted(unknown)}")
        if "status" in fields and fields["status"] not in CASE_STATUSES:
            raise CaseError(f"Status must be one of {CASE_STATUSES}")
        if fields.get("status") == "approved" and not (fields.get("reviewer") or "").strip():
            raise CaseError("Approving a case requires the name of the human reviewer")
        with _lock:
            wb = self._open()
            ws = wb["Cases"]
            row = self._find_row(ws, lambda r: r["case_id"] == case_id)
            if row is None:
                raise CaseError(f"Case {case_id} not found")
            for key, value in fields.items():
                if value is not None:
                    self._set(ws, row, key, value)
            self._set(ws, row, "updated_at", _now())
            self._log(wb, case_id, actor, "updated", json.dumps({k: v for k, v in fields.items() if v is not None}))
            self._save(wb)
        return self.get_case(case_id)

    # ---------------- documents ----------------
    def _append_document(self, wb, case_id: str, doc: dict) -> str:
        ws = wb["Documents"]
        existing = [d["doc_id"] for d in self._rows(ws) if d["case_id"] == case_id]
        doc_id = doc.get("id") or f"doc-{len(existing) + 1}"
        if doc_id in existing:
            doc_id = f"{doc_id}-{len(existing) + 1}"
        row = {
            "case_id": case_id,
            "doc_id": doc_id,
            "name": doc.get("name", ""),
            "kind": doc.get("kind", "legal"),
            "mandatory": bool(doc.get("mandatory", False)),
            "basis": json.dumps(doc.get("basis", []), ensure_ascii=False),
            "status": "missing",
            "due_date": doc.get("due_date", ""),
            "file": "",
            "updated_at": _now(),
            "note": doc.get("why", ""),
        }
        ws.append([row[c] for c in DOC_COLUMNS])
        return doc_id

    def add_document(self, case_id: str, name: str, mandatory: bool, due_date: str = "", actor: str = "") -> str:
        self.get_case(case_id)
        with _lock:
            wb = self._open()
            doc_id = self._append_document(
                wb, case_id, {"id": None, "name": name, "mandatory": mandatory, "due_date": due_date, "kind": "custom"}
            )
            self._log(wb, case_id, actor, "document_added", name)
            self._save(wb)
        return doc_id

    def update_document(self, case_id: str, doc_id: str, actor: str = "", **fields) -> dict:
        _check_case_id(case_id)
        allowed = {"status", "due_date", "note", "file"}
        if set(fields) - allowed:
            raise CaseError(f"Cannot update fields: {sorted(set(fields) - allowed)}")
        if "status" in fields and fields["status"] not in DOC_STATUSES:
            raise CaseError(f"Document status must be one of {DOC_STATUSES}")
        if fields.get("due_date"):
            date.fromisoformat(str(fields["due_date"]))
        with _lock:
            wb = self._open()
            ws = wb["Documents"]
            row = self._find_row(ws, lambda r: r["case_id"] == case_id and r["doc_id"] == doc_id)
            if row is None:
                raise CaseError(f"Document {doc_id} not found in {case_id}")
            for key, value in fields.items():
                if value is not None:
                    self._set(ws, row, key, value)
            self._set(ws, row, "updated_at", _now())
            self._log(wb, case_id, actor, "document_updated", f"{doc_id}: {json.dumps(fields)}")
            self._save(wb)
        return self.get_case(case_id)

    def save_upload(self, case_id: str, doc_id: str, filename: str, content: bytes, actor: str = "") -> str:
        _check_case_id(case_id)
        safe = _SAFE_NAME.sub("_", Path(filename).name).strip("._") or "upload"
        folder = self.files_dir / case_id
        folder.mkdir(parents=True, exist_ok=True)
        target = folder / f"{_SAFE_NAME.sub('_', doc_id)}__{safe}"
        target.write_bytes(content)
        relative = target.relative_to(self.files_dir).as_posix()
        self.update_document(case_id, doc_id, actor=actor, status="received", file=relative)
        return relative


def _truthy(value) -> bool:
    return value is True or str(value).strip().lower() in ("true", "1", "yes")


def _json_or_empty(value):
    try:
        return json.loads(value) if value else []
    except (TypeError, json.JSONDecodeError):
        return []


def _alerts(docs: list[dict]) -> list[str]:
    alerts = []
    today = date.today()
    for d in docs:
        if d["status"] != "missing":
            continue
        if d["mandatory"]:
            alerts.append(f"Mandatory document missing: {d['name']}")
        due = str(d.get("due_date") or "")[:10]
        if due:
            try:
                if date.fromisoformat(due) < today:
                    alerts.append(f"Overdue since {due}: {d['name']}")
            except ValueError:
                pass
    return alerts
